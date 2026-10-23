import unittest
from io import BytesIO
from openpyxl import load_workbook
from app import app

class QuotationTests(unittest.TestCase):
    def setUp(self):
        self.client=app.test_client()
        self.products=self.client.get('/api/catalogue').get_json()
    def payload(self,p):
        return dict(product=p,environment=p['environment'],unit='m',wall_width=6,wall_height=4,mode='grid',columns=5,rows=4,controller='HD-KV410',quotation={})
    def test_variants_and_quote(self):
        variants=[p for p in self.products if p['id'].startswith('pm-')]
        self.assertEqual(len(variants),5)
        self.assertEqual({p['technology'] for p in variants},{'COB','SMD'})
        for p in variants:
            response=self.client.post('/api/quotation',json=self.payload(p))
            self.assertEqual(response.status_code,200,response.get_data(as_text=True) if response.status_code!=200 else '')
            wb=load_workbook(BytesIO(response.data));ws=wb['Quotation']
            self.assertIn(p['name'],ws['C10'].value)
            self.assertIn('20 cabinets',ws['C10'].value)
            self.assertIsNone(ws['F10'].value)
            self.assertIn('Rates pending',ws['J14'].value)
            self.assertEqual(len(wb['Design & Specifications']._images),1)
    def test_rates_formulas_and_text_safety(self):
        d=self.payload(self.products[-1]);d['quotation']={'client':'=HYPERLINK("bad")','salesperson':'Sales','contact':'555-0100','wall_rate':10000,'controller_rate':500,'mount_rate':200,'installation_rate':100,'gst':18,'controller_qty':2,'controller_description':'=1+1'}
        r=self.client.post('/api/quotation',json=d);self.assertEqual(r.status_code,200)
        ws=load_workbook(BytesIO(r.data))['Quotation']
        self.assertTrue(str(ws['F10'].value).startswith('=IF('));self.assertEqual(ws['G10'].value,0.18);self.assertEqual(ws['D11'].value,2)
        self.assertIsNone(ws['F11'].value)
        self.assertEqual(ws['J11'].value,'=IF(F11="","",ROUND(D11*I11,2))')
        self.assertIn('Rates pending',ws['J16'].value)
        self.assertEqual(load_workbook(BytesIO(r.data),data_only=True)['Quotation']['J16'].value,'Rates pending')
        specs=load_workbook(BytesIO(r.data))['Design & Specifications']
        controller_pricing=next(row[1].value for row in specs.iter_rows(min_col=1,max_col=2) if row[0].value=='Controller pricing')
        self.assertEqual(controller_pricing,'Pending revised controller rates; not calculated.')
        self.assertNotEqual(ws['A1'].data_type,'f');self.assertNotEqual(ws['C11'].data_type,'f')
    def test_invalid_rates_rejected(self):
        for bad in [-1,'NaN','Infinity','abc']:
            d=self.payload(self.products[0]);d['quotation']['wall_rate']=bad
            self.assertEqual(self.client.post('/api/quotation',json=d).status_code,400)
    def test_variant_geometry(self):
        for p in self.products:
            if p.get('technology'):
                result=self.client.post('/api/calculate',json=self.payload(p)).get_json()
                self.assertEqual(result['pixels_w'],None if p['pixels_w'] is None else p['pixels_w']*5)
                self.assertAlmostEqual(result['width_m'],p['width_mm']*5/1000)
