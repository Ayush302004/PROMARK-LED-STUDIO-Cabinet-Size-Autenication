"""Verify that extra Excel rows and long descriptions never create a second page."""
import json
import re
import unittest
from io import BytesIO

from openpyxl import load_workbook

from test_mounting_quotation import BASE, pdf_content
from quotation import TERMS, build_pdf_from_excel, build_quotation, build_quotation_pdf, read_excel_quotation


def quotation_data():
    return {
        'product': json.loads((BASE / 'catalogue.json').read_text(encoding='utf-8'))[0],
        'controller': 'HD-VP620A', 'unit': 'ft', 'mode': 'dimensions',
        'wall_width': 30, 'wall_height': 20, 'screen_width': 20.997, 'screen_height': 9.449,
        'quotation': {'client': 'Page fit check', 'contact': 'test@example.com',
                      'salesperson': 'Test', 'wall_rate': 1000, 'mount_type': 'unipole', 'gst': 18},
    }


def edited_workbook(extra_rows, description_repeats=1, extra_terms=0):
    workbook = load_workbook(build_quotation(quotation_data(), BASE), data_only=True)
    sheet = workbook['Quotation']
    # openpyxl does not shift merge ranges when inserting rows; keep this fixture
    # unmerged so it preserves the same values as an actual Excel row insertion.
    for merged in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged))
    if extra_rows:
        sheet.insert_rows(14, extra_rows)
    for index in range(extra_rows):
        description = ('Additional equipment, cabling and installation details. ' * description_repeats
                       + f'ENDITEM{index:02d}')
        for column, value in enumerate([index + 5, f'Added item {index + 1}', description,
                                       2, 'Nos', 1250, 0.18], 1):
            sheet.cell(14 + index, column, value)
    if extra_terms:
        sheet.insert_rows(28 + extra_rows, extra_terms)
        for index in range(extra_terms):
            sheet.cell(28 + extra_rows + index, 1, f'{10 + index}.) Additional condition {index + 1}.')
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    output.seek(0)
    return output


class QuotationPageFitTests(unittest.TestCase):
    def assert_single_page(self, output):
        self.assertEqual(len(re.findall(rb'/Type\s*/Page\b', output.getvalue())), 1)
        content = pdf_content(output)
        self.assertIn('Total Amount', content)
        self.assertIn('GST Amount', content)
        self.assertIn('LAN ports for AV System', content)
        self.assertIn('Cable and speaker quantities', content)
        self.assertEqual(content.count('PROMARK TECHNOLOGIES PRIVATE LIMITED'), 1)
        return content

    def test_direct_pdf_keeps_final_terms_on_first_page(self):
        self.assert_single_page(build_quotation_pdf(quotation_data(), BASE))

    def test_excel_with_added_rows_and_terms_fits_one_page(self):
        for count, repeats in ((1, 1), (5, 2), (15, 3), (30, 5)):
            with self.subTest(added_rows=count, description_repeats=repeats):
                source = edited_workbook(count, repeats, extra_terms=3)
                workbook = load_workbook(source, data_only=True)
                _, rows, terms, total, gst = read_excel_quotation(workbook['Quotation'])
                self.assertEqual(len(rows), 4 + count)
                self.assertEqual(len(terms), len(TERMS) + 3)
                workbook.close()
                source.seek(0)
                content = self.assert_single_page(build_pdf_from_excel(source))
                for index in range(count):
                    self.assertIn(f'Added item {index + 1}', content)
                    self.assertIn(f'ENDITEM{index:02d}', content)
                for index in range(3):
                    self.assertIn(f'Additional condition {index + 1}', content)
                self.assertIn(f'{total:,.2f}', content)
                self.assertIn(f'{gst:,.2f}', content)

    def test_long_existing_row_fits_without_splitting(self):
        source = edited_workbook(0)
        workbook = load_workbook(source, data_only=True)
        workbook['Quotation']['C10'] = 'Expanded specification details. ' * 120 + 'FINALDETAIL'
        output = BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)
        content = self.assert_single_page(build_pdf_from_excel(output))
        self.assertIn('FINALDETAIL', content)


if __name__ == '__main__':
    unittest.main()
