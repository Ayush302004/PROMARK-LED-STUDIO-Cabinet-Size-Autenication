"""Run with: python -m unittest discover -s tests."""
import base64
import json
import re
import sys
import unittest
import zlib
from decimal import Decimal
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parents[1] / 'av_wall_designer'
sys.path.insert(0, str(BASE))
from quotation import build_pdf_from_excel, build_quotation, build_quotation_pdf, read_excel_quotation


def pdf_content(output):
    """Read ReportLab's page drawing commands without an extra PDF dependency."""
    streams = []
    for stream in re.findall(rb'stream\r?\n(.*?)endstream', output.getvalue(), re.S):
        try:
            streams.append(zlib.decompress(base64.a85decode(stream.strip(), adobe=True)))
        except (ValueError, zlib.error):
            pass
    return b'\n'.join(streams).decode('latin-1')


class MountingQuotationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.product = json.loads((BASE / 'catalogue.json').read_text(encoding='utf-8'))[0]

    def data(self, **quotation):
        return {
            'product': self.product, 'controller': 'HD-KV410',
            'unit': 'ft', 'mode': 'dimensions', 'preview_mode': 'exact',
            'wall_width': 20, 'wall_height': 20, 'screen_width': 10, 'screen_height': 10,
            'quotation': {'client': 'Test client', 'salesperson': 'Test', 'contact': 'test@example.com',
                          'wall_rate': 1000, 'gst': 18, **quotation},
        }

    def check_exports(self, data, expected_name, expected_price):
        output = build_quotation(data, BASE)
        workbook = load_workbook(output, data_only=True)
        sheet = workbook['Quotation']
        _, rows, _, total, _ = read_excel_quotation(sheet)
        mounting = [row for row in rows if row[1] in ('Uni-pole', 'Mounting structure')]
        if expected_name is None:
            self.assertEqual(mounting, [])
            self.assertEqual(sheet['B12'].value, 'Installation')
            self.assertEqual(sheet['J13'].value, float(total))
            self.assertIn('F12', workbook['Design & Specifications']['A5'].value)
        else:
            self.assertEqual(len(mounting), 1)
            self.assertEqual(mounting[0][1], expected_name)
            self.assertEqual(sheet['D12'].value, 1)
            self.assertEqual(Decimal(str(sheet['F12'].value)), expected_price)
            tax = (expected_price * Decimal('0.18')).quantize(Decimal('0.01'), rounding='ROUND_HALF_UP')
            self.assertEqual(Decimal(str(sheet['H12'].value)), tax)
            self.assertEqual(Decimal(str(sheet['J12'].value)), expected_price + tax)
            self.assertEqual(sheet['B13'].value, 'Installation')
            self.assertEqual(sheet['J14'].value, float(total))
        self.assertEqual([row[0] for row in rows], [str(i) for i in range(1, len(rows) + 1)])
        workbook.close()
        direct_pdf = build_quotation_pdf(data, BASE)
        self.assertEqual(len(re.findall(rb'/Type\s*/Page\b', direct_pdf.getvalue())), 1)
        direct = pdf_content(direct_pdf)
        output.seek(0)
        reuploaded = build_pdf_from_excel(output)
        self.assertEqual(len(re.findall(rb'/Type\s*/Page\b', reuploaded.getvalue())), 1)
        for content in (direct, pdf_content(reuploaded)):
            self.assertIn('Installation', content)
            for name in ('Uni-pole', 'Mounting structure'):
                self.assertEqual(name in content, name == expected_name)
            if expected_name:
                self.assertIn(f'{expected_price:,.2f}', content)
        return output

    def test_none_removes_mounting_even_with_stale_price(self):
        self.check_exports(self.data(mount_type='none', mount_rate=175000, mount_qty=100), None, None)

    def test_missing_selection_defaults_to_none(self):
        self.check_exports(self.data(), None, None)

    def test_unipole_has_one_fixed_unit_price(self):
        self.check_exports(self.data(mount_type='unipole', mount_qty=100), 'Uni-pole', Decimal('175000'))

    def test_structure_uses_screen_area_and_editable_excel_formula(self):
        output = self.check_exports(self.data(mount_type='structure', mount_qty=999), 'Mounting structure', Decimal('80000'))
        output.seek(0)
        workbook = load_workbook(output)
        self.assertEqual(workbook['Quotation']['F12'].value, "=ROUND('Design & Specifications'!B4*'Design & Specifications'!B3,2)")
        self.assertEqual(workbook['Design & Specifications']['B4'].value, 800)
        workbook.close()

    def test_custom_rate_and_description(self):
        data = self.data(mount_type='structure', mount_rate=900, mount_description='Custom mounting <frame> & supports.')
        self.check_exports(data, 'Mounting structure', Decimal('90000'))

    def test_zero_rate_is_a_priced_mounting_item(self):
        self.check_exports(self.data(mount_type='unipole', mount_rate=0), 'Uni-pole', Decimal('0'))

    def test_structure_metres_and_fractional_area(self):
        data = self.data(mount_type='structure')
        data.update(unit='m', screen_width=3, screen_height=2)
        expected = (Decimal('800') * Decimal('6') / Decimal('0.09290304')).quantize(Decimal('0.01'), rounding='ROUND_HALF_UP')
        self.check_exports(data, 'Mounting structure', expected)

    def test_structure_uses_buildable_area_when_selected(self):
        data = self.data(mount_type='structure')
        data.update(unit='m', preview_mode='buildable', screen_width=3, screen_height=2)
        # Five 0.6 m cabinets by five 0.3375 m cabinets: displayed 3.000 x 1.688 m.
        expected = (Decimal('800') * Decimal('3.000') * Decimal('1.688') / Decimal('0.09290304')).quantize(Decimal('0.01'), rounding='ROUND_HALF_UP')
        self.check_exports(data, 'Mounting structure', expected)

    def test_invalid_mounting_values_are_rejected(self):
        for choice in ({'mount_type': 'invalid'}, {'mount_type': 'structure', 'mount_rate': -1},
                       {'mount_type': 'unipole', 'mount_rate': 'nan'}):
            for build in (build_quotation, build_quotation_pdf):
                with self.subTest(choice=choice, format=build.__name__), self.assertRaises(ValueError):
                    build(self.data(**choice), BASE)


if __name__ == '__main__':
    unittest.main()
